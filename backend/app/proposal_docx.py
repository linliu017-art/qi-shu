"""将方案转换成可读的 Word；不调用模型，不下载外部资源。"""
import io
import re
from collections import defaultdict
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from markdown_it import MarkdownIt

UUID = re.compile(r'\b[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}\b', re.I)

def clean_text(text, names):
    return UUID.sub(lambda m: names.get(m.group(0).lower(), '相关信息待核对'), text)

def inline_text(token):
    if not token.children:
        return token.content
    result = []
    for child in token.children:
        if child.type in ('text', 'code_inline', 'html_inline'):
            result.append(child.content)
        elif child.type in ('softbreak', 'hardbreak'):
            result.append('\n')
        elif child.type == 'image':
            result.append(child.content or '图片（需核对原文）')
    return ''.join(result)

def add_inline(paragraph, token, names):
    bold = italic = False
    for child in token.children or []:
        if child.type == 'strong_open': bold = True
        elif child.type == 'strong_close': bold = False
        elif child.type == 'em_open': italic = True
        elif child.type == 'em_close': italic = False
        elif child.type in ('text', 'code_inline', 'html_inline'):
            run = paragraph.add_run(clean_text(child.content, names)); run.bold = bold; run.italic = italic
        elif child.type in ('softbreak', 'hardbreak'): paragraph.add_run('\n')
        elif child.type == 'image': paragraph.add_run(child.content or '图片（需核对原文）')

def split_sources_table(rows, names):
    if not rows: return rows
    headers = [v.strip() for v in rows[0]]
    missing = next((i for i, value in enumerate(headers) if value in ('需补充资料','需要补充资料','还需补充')), -1)
    if missing < 0 or not any(missing < len(row) and UUID.search(row[missing]) for row in rows[1:]):
        return [[clean_text(c, names) for c in row] for row in rows]
    source_index = next((i for i, value in enumerate(headers) if value in ('已有资料','已有资料出处')), -1)
    result = []
    for index, row in enumerate(rows):
        row = list(row)
        while len(row) < len(headers): row.append('')
        if index == 0:
            row[missing] = '还需补充'
            if source_index < 0: row.insert(missing, '已有资料')
        else:
            ids = UUID.findall(row[missing])
            listed = list(dict.fromkeys(names.get(id.lower(), '相关资料待核对') for id in ids))
            if ids:
                remaining = UUID.sub('', row[missing]).strip(' ，,；;、')
                row[missing] = remaining or '待确认'
            if source_index < 0:
                row.insert(missing, '；'.join(listed) if listed else '未列出')
            elif listed:
                row[source_index] = '；'.join(filter(None, [row[source_index], *listed]))
        result.append([clean_text(c, names) for c in row])
    return result

def build_docx(content, kind, version, questions, material_names, citations, baseline_description):
    doc = Document()
    tokens = MarkdownIt('commonmark', {'html': False}).enable('table').parse(content)
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.5)
    normal = doc.styles['Normal']
    normal.font.name = '宋体'; normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    doc.add_heading('方案草稿（未审核）' if kind == 'draft' else '已审核方案（本机试用）', 0)
    doc.add_paragraph(f'方案版本：第 {version} 版')
    doc.add_paragraph('依据的需求清单：' + baseline_description)
    doc.add_paragraph('审核状态：草稿，未审核' if kind == 'draft' else '审核状态：已审核（本机试用）')
    doc.add_heading('目录', 1)
    headings = [(int(token.tag[1:]), inline_text(tokens[index + 1])) for index, token in enumerate(tokens[:-1]) if token.type == 'heading_open' and tokens[index + 1].type == 'inline' and int(token.tag[1:]) <= 3]
    # Keep a readable initial list inside the Word field; office software can update it.
    toc_paragraph = doc.add_paragraph()
    toc = OxmlElement('w:fldSimple')
    toc.set(qn('w:instr'), 'TOC \\o "1-3" \\h \\z \\u')
    toc_titles = [clean_text(title, material_names) for _, title in headings]
    if questions: toc_titles.append('需求清单中的问题处理情况' if kind == 'approved' else '需求清单中的未解决事项')
    if citations: toc_titles.append('资料来源')
    for index, title in enumerate(toc_titles):
        run = OxmlElement('w:r')
        if index: run.append(OxmlElement('w:br'))
        text = OxmlElement('w:t'); text.text = title; run.append(text); toc.append(run)
    toc_paragraph._p.append(toc)
    settings = doc.settings.element
    update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true'); settings.append(update)
    footer = section.footer.paragraphs[0]
    footer.add_run('第 ')
    page = OxmlElement('w:fldSimple'); page.set(qn('w:instr'), 'PAGE')
    page_run = OxmlElement('w:r'); page_text = OxmlElement('w:t'); page_text.text = '1'; page_run.append(page_text); page.append(page_run); footer._p.append(page)
    footer.add_run(' 页')
    i = 0; list_styles = []; metadata_level = None
    while i < len(tokens):
        token = tokens[i]
        if token.type in ('bullet_list_open', 'ordered_list_open'):
            list_styles.append('List Bullet' if token.type == 'bullet_list_open' else 'List Number')
        elif token.type in ('bullet_list_close', 'ordered_list_close'):
            if list_styles: list_styles.pop()
        elif token.type == 'heading_open' and i+1 < len(tokens) and tokens[i+1].type == 'inline':
            level = max(1, min(6, int(token.tag[1:])))
            heading = inline_text(tokens[i+1]).strip()
            if metadata_level is not None and level <= metadata_level: metadata_level = None
            if re.fullmatch(r'(?:文档|文件|方案)(?:基本)?信息', heading): metadata_level = level
            paragraph = doc.add_heading(level=level); add_inline(paragraph, tokens[i+1], material_names); i += 1
        elif token.type == 'paragraph_open' and i+1 < len(tokens) and tokens[i+1].type == 'inline':
            paragraph = doc.add_paragraph(style=list_styles[-1] if list_styles else None)
            original = inline_text(tokens[i+1])
            if kind == 'approved' and metadata_level is not None and re.search(r'(文档状态|审核状态|尚未审核|未审核|待审核|方案初稿|文档版本)', original):
                paragraph.add_run('原稿生成时记录（当前状态以文首为准）：')
            add_inline(paragraph, tokens[i+1], material_names); i += 1
        elif token.type in ('fence', 'code_block'):
            doc.add_paragraph(clean_text(token.content, material_names))
        elif token.type == 'table_open':
            rows = []; row = None; cell = None; i += 1
            while i < len(tokens) and tokens[i].type != 'table_close':
                current = tokens[i]
                if current.type == 'tr_open': row = []
                elif current.type in ('th_open','td_open'): cell = ''
                elif current.type == 'inline' and cell is not None: cell += inline_text(current)
                elif current.type in ('th_close','td_close') and row is not None:
                    row.append(cell or ''); cell = None
                elif current.type == 'tr_close' and row is not None:
                    rows.append(row); row = None
                i += 1
            rows = split_sources_table(rows, material_names)
            if rows:
                count = max(len(row) for row in rows)
                table = doc.add_table(rows=0, cols=count); table.style = 'Table Grid'
                borders = OxmlElement('w:tblBorders')
                for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                    border = OxmlElement(f'w:{edge}'); border.set(qn('w:val'), 'single'); border.set(qn('w:sz'), '4'); border.set(qn('w:color'), 'AABBAE'); borders.append(border)
                table._tbl.tblPr.append(borders)
                for rindex, values in enumerate(rows):
                    cells = table.add_row().cells
                    for j, value in enumerate(values):
                        cells[j].text = value
                        if rindex == 0:
                            for run in cells[j].paragraphs[0].runs: run.bold = True
        i += 1
    if questions:
        doc.add_heading('需求清单中的问题处理情况' if kind == 'approved' else '需求清单中的未解决事项', 1)
        for question in questions:
            text = question.get('text', '')
            answer = question.get('answer', '')
            review = question.get('review') or {}
            label = {'resolved': '已解决', 'noncommit': '已明确不作承诺', 'general_approved': '一般事项批准保留（未解决）'}.get(review.get('disposition'), '')
            suffix = ('；处理结果：' + label + '；说明：' + clean_text(review.get('reason', ''), material_names)) if label else ''
            doc.add_paragraph(clean_text(text, material_names) + (('；已有答复：' + clean_text(answer, material_names)) if answer else '') + suffix, style='List Bullet')
    if citations:
        doc.add_heading('资料来源', 1)
        grouped = defaultdict(list)
        for name, quote in citations:
            if quote and quote not in grouped[name]: grouped[name].append(quote)
        for name, quotes in grouped.items():
            doc.add_paragraph(name if re.search(r'第\s*\d+\s*版', name) else name + '（资料版本未记录）', style='List Bullet')
            for quote in quotes: doc.add_paragraph('原文摘录：' + quote)
    stream = io.BytesIO(); doc.save(stream)
    return stream.getvalue()
