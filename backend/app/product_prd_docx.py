"""Render saved PRD Markdown into a simple, editable Word document."""
from io import BytesIO
from docx import Document
from markdown_it import MarkdownIt

def plain_inline(token):
    result=[]
    for item in token.children or []:
        if item.type in ('text','code_inline','html_inline'):
            result.append(item.content)
        elif item.type in ('softbreak','hardbreak'):
            result.append('\n')
        elif item.type=='image':
            result.append(item.content or '[图片]')
    return ''.join(result)

def render_prd(title:str, content:str)->bytes:
    document=Document()
    document.add_heading(title,0)
    tokens=MarkdownIt('commonmark').enable('table').parse(content)
    if len(tokens)>20000: raise ValueError('文档内容过于复杂。')
    heading=0; lists=[]; table=None; row=None; cell=None
    for token in tokens:
        kind=token.type
        if kind=='heading_open': heading=int(token.tag[1])
        elif kind=='heading_close': heading=0
        elif kind=='bullet_list_open': lists.append('List Bullet')
        elif kind=='ordered_list_open': lists.append('List Number')
        elif kind in ('bullet_list_close','ordered_list_close') and lists: lists.pop()
        elif kind=='table_open': table=[]
        elif kind=='tr_open' and table is not None: row=[]
        elif kind in ('th_open','td_open') and row is not None: cell=[]
        elif kind=='inline':
            text=plain_inline(token)
            if cell is not None: cell.append(text)
            elif heading: document.add_heading(text,level=min(heading,9))
            else: document.add_paragraph(text,style=lists[-1] if lists else None)
        elif kind in ('th_close','td_close') and row is not None and cell is not None:
            row.append(' '.join(cell));cell=None
        elif kind=='tr_close' and table is not None and row is not None:
            table.append(row);row=None
        elif kind=='table_close' and table is not None:
            width=max((len(r) for r in table),default=0)
            if len(table)>200 or width>30: raise ValueError('表格过大。')
            if width:
                output=document.add_table(rows=len(table),cols=width)
                output.style='Table Grid'
                for i,values in enumerate(table):
                    for j,value in enumerate(values): output.cell(i,j).text=value
            table=None
        elif kind in ('fence','code_block'):
            document.add_paragraph(token.content)
    out=BytesIO();document.save(out);return out.getvalue()
