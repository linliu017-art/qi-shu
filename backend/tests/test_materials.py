import importlib
import pytest
from fastapi.testclient import TestClient

@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv('QISHU_DATA_DIR', str(tmp_path))
    monkeypatch.setenv('DEEPSEEK_API_KEY', '')
    import app.database as database
    import app.models as models
    import app.main as main
    importlib.reload(database)
    importlib.reload(models)
    importlib.reload(main)
    with TestClient(main.app) as c:
        yield c
    database.engine.dispose()

def test_create_project_and_restore_material(client):
    response = client.post('/api/v1/projects', json={'name': '虚构项目'})
    assert response.status_code == 201
    project_id = response.json()['id']
    response = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('访问.txt', '客户希望建立台账。'.encode(), 'text/plain')})
    assert response.status_code == 201
    assert response.json()['text'] == '客户希望建立台账。'
    assert client.get(f'/api/v1/projects/{project_id}/materials').json()[0]['text'] == '客户希望建立台账。'

def test_reject_invalid_inputs_and_missing_project(client):
    assert client.post('/api/v1/projects', json={'name': '  '}).status_code == 422
    missing = client.post('/api/v1/projects/missing/materials', files={'file': ('a.txt', b'abc', 'text/plain')})
    assert missing.status_code == 404
    project_id = client.post('/api/v1/projects', json={'name': '项目'}).json()['id']
    invalid = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('a.txt', b'\x00fake', 'text/plain')})
    assert invalid.status_code == 422
    assert invalid.json()['error']['code'] == 'INVALID_FILE'
    pdf = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('a.pdf', b'%PDF-', 'application/pdf')})
    assert pdf.json()['error']['code'] == 'INVALID_FILE'

def test_mock_model_review_and_baseline(client, monkeypatch):
    import json
    import app.analysis as analysis
    import app.database as database
    monkeypatch.setattr(analysis, 'SessionLocal', database.SessionLocal)
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'fake-for-test')
    project_id = client.post('/api/v1/projects', json={'name': '样例'}).json()['id']
    material_id = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('a.txt', '客户希望建立台账。'.encode(), 'text/plain')}).json()['id']
    class Response:
        status_code = 200
        def raise_for_status(self): pass
        def json(self):
            return {'choices': [{'message': {'content': json.dumps({'requirements': [{'text': '建立台账', 'material_id': material_id, 'quote': '客户希望建立台账。'}]})}}]}
    monkeypatch.setattr(analysis.httpx, 'post', lambda *a, **kw: Response())
    task = client.post(f'/api/v1/projects/{project_id}/requirement-tasks', json={'material_ids': [material_id], 'confirm_external_use': True})
    assert task.status_code == 202
    result = client.get(f"/api/v1/projects/{project_id}/requirement-tasks/{task.json()['id']}").json()
    assert result['status'] == 'completed'
    item = client.get(f'/api/v1/projects/{project_id}/requirements').json()[0]
    assert item['status'] == 'pending'
    assert client.post(f'/api/v1/projects/{project_id}/baselines').status_code == 422
    review = client.patch(f"/api/v1/projects/{project_id}/requirements/{item['id']}", json={'status': 'confirmed', 'text': item['text'], 'version': item['version']})
    assert review.status_code == 200
    assert client.post(f'/api/v1/projects/{project_id}/baselines').json()['count'] == 1
    baseline = client.get(f'/api/v1/projects/{project_id}/baselines').json()[0]
    assert baseline['requirements'][0]['text'] == '建立台账'
    assert client.patch(f"/api/v1/projects/{project_id}/requirements/{item['id']}", json={'status': 'rejected', 'text': item['text'], 'version': item['version']}).status_code == 409
    path = f"/api/v1/projects/{project_id}/requirements/{item['id']}"
    confirmed = review.json()
    repeated = client.patch(path, json={'status': 'confirmed', 'text': confirmed['text'], 'version': confirmed['version']}).json()
    assert repeated['version'] == confirmed['version']
    undone = client.patch(path, json={'status': 'pending', 'text': confirmed['text'], 'version': confirmed['version']}).json()
    assert undone['status'] == 'pending' and undone['version'] == confirmed['version'] + 1
    assert client.post(f'/api/v1/projects/{project_id}/baselines').status_code == 422
    assert client.get(f'/api/v1/projects/{project_id}/baselines').json()[0]['requirements'][0]['text'] == '建立台账'

def test_mock_model_invalid_quote_fails_without_suggestions(client, monkeypatch):
    import json
    import app.analysis as analysis
    import app.database as database
    monkeypatch.setattr(analysis, 'SessionLocal', database.SessionLocal)
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'fake-for-test')
    project_id = client.post('/api/v1/projects', json={'name': '样例'}).json()['id']
    material_id = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('a.txt', b'actual source', 'text/plain')}).json()['id']
    class Response:
        status_code = 200
        def raise_for_status(self): pass
        def json(self):
            return {'choices': [{'message': {'content': json.dumps({'requirements': [{'text': 'X', 'material_id': material_id, 'quote': 'made up'}]})}}]}
    monkeypatch.setattr(analysis.httpx, 'post', lambda *a, **kw: Response())
    task = client.post(f'/api/v1/projects/{project_id}/requirement-tasks', json={'material_ids': [material_id], 'confirm_external_use': True}).json()
    assert client.get(f"/api/v1/projects/{project_id}/requirement-tasks/{task['id']}").json()['status'] == 'failed'
    assert client.get(f'/api/v1/projects/{project_id}/requirements').json() == []

def test_original_is_saved_and_scoped_to_project(client):
    project_id = client.post('/api/v1/projects', json={'name': '项目一'}).json()['id']
    other_id = client.post('/api/v1/projects', json={'name': '项目二'}).json()['id']
    content = '虚构客户资料'.encode()
    saved = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('a.txt', content, 'text/plain')}).json()
    assert client.get(f"/api/v1/projects/{project_id}/materials/{saved['id']}/original").content == content
    assert client.get(f"/api/v1/projects/{other_id}/materials/{saved['id']}/original").status_code == 404

def test_segment_limit_and_resume_preserve_progress(client, monkeypatch):
    import json
    import app.analysis as analysis
    import app.database as database
    monkeypatch.setattr(analysis, 'SessionLocal', database.SessionLocal)
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'fake-for-test')
    monkeypatch.setenv('QISHU_MAX_REQUESTS', '1')
    project_id = client.post('/api/v1/projects', json={'name': '分段'}).json()['id']
    content = '甲' * 11999 + '\n' + '乙' * 1000
    material_id = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('a.txt', content.encode(), 'text/plain')}).json()['id']
    class Response:
        status_code = 200
        def raise_for_status(self): pass
        def json(self):
            return {'usage': {'prompt_tokens': 12, 'completion_tokens': 4}, 'choices': [{'message': {'content': json.dumps({'requirements': []})}}]}
    monkeypatch.setattr(analysis.httpx, 'post', lambda *a, **kw: Response())
    task = client.post(f'/api/v1/projects/{project_id}/requirement-tasks', json={'material_ids': [material_id], 'confirm_external_use': True}).json()
    url = f"/api/v1/projects/{project_id}/requirement-tasks/{task['id']}"
    paused = client.get(url).json()
    assert paused['status'] == 'paused_limit'
    assert paused['segments_total'] == 2 and paused['segments_completed'] == 1
    assert paused['calls'] == 1 and paused['prompt_tokens'] == 12
    assert client.post(url + '/resume', json={'confirm_external_use': True}).json()['error']['code'] == 'CALL_LIMIT'
    monkeypatch.setenv('QISHU_MAX_REQUESTS', '2')
    assert client.post(url + '/resume', json={'confirm_external_use': True}).status_code == 200
    finished = client.get(url).json()
    assert finished['status'] == 'completed' and finished['segments_completed'] == 2
    assert finished['calls'] == 2 and finished['prompt_tokens'] == 24

def test_resume_retains_reviewed_results_and_unknown_usage(client, monkeypatch):
    import json
    import app.analysis as analysis
    import app.database as database
    monkeypatch.setattr(analysis, 'SessionLocal', database.SessionLocal)
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'fake-for-test')
    project_id = client.post('/api/v1/projects', json={'name': '中断恢复'}).json()['id']
    content = '甲' * 11999 + '\n' + '乙' * 1000
    material_id = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('a.txt', content.encode(), 'text/plain')}).json()['id']
    requests = []
    class Response:
        def __init__(self, fail): self.status_code = 400 if fail else 200
        def raise_for_status(self): pass
        def json(self): return {'choices': [{'message': {'content': json.dumps({'requirements': [{'text': '记录需求', 'material_id': material_id, 'quote': '甲甲'}]})}}]}
    def fake_post(*args, **kwargs):
        requests.append(1)
        if len(requests) == 2: return Response(True)
        if len(requests) >= 3:
            class EmptyResponse(Response):
                def json(self): return {'usage': {'prompt_tokens': 4, 'completion_tokens': 2}, 'choices': [{'message': {'content': json.dumps({'requirements': []})}}]}
            return EmptyResponse(False)
        return Response(False)
    monkeypatch.setattr(analysis.httpx, 'post', fake_post)
    task = client.post(f'/api/v1/projects/{project_id}/requirement-tasks', json={'material_ids': [material_id], 'confirm_external_use': True}).json()
    url = f"/api/v1/projects/{project_id}/requirement-tasks/{task['id']}"
    interrupted = client.get(url).json()
    assert interrupted['status'] == 'failed' and interrupted['segments_completed'] == 1
    requirement = client.get(f'/api/v1/projects/{project_id}/requirements').json()[0]
    path = f"/api/v1/projects/{project_id}/requirements/{requirement['id']}"
    assert client.patch(path, json={'status': 'confirmed', 'text': '人工确认的内容', 'version': 1}).status_code == 200
    assert client.post(url + '/resume', json={'confirm_external_use': True}).status_code == 200
    finished = client.get(url).json()
    assert finished['status'] == 'completed' and finished['segments_completed'] == 2
    assert finished['calls'] == 3 and finished['usage_complete'] is False and finished['estimated_cost_usd'] is None
    rows = client.get(f'/api/v1/projects/{project_id}/requirements').json()
    assert len(rows) == 1 and rows[0]['text'] == '人工确认的内容' and rows[0]['status'] == 'confirmed'

def test_key_is_not_returned_and_model_must_be_allowed(client, monkeypatch):
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'fake-for-test')
    config = client.get('/api/v1/config').json()
    assert config['model_ready'] is True and 'key' not in str(config).lower()
    project_id = client.post('/api/v1/projects', json={'name': '模型验证'}).json()['id']
    material_id = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('a.txt', b'test', 'text/plain')}).json()['id']
    response = client.post(f'/api/v1/projects/{project_id}/requirement-tasks', json={'material_ids': [material_id], 'model_name': 'unapproved-model', 'confirm_external_use': True})
    assert response.status_code == 422 and response.json()['error']['code'] == 'MODEL_NOT_ALLOWED'

def test_restart_marks_unfinished_task_and_keeps_source(client):
    import json
    import app.database as database
    import app.models as models
    import app.main as main
    project_id = client.post('/api/v1/projects', json={'name': '中断测试'}).json()['id']
    material = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('source.txt', b'example', 'text/plain')}).json()
    with database.SessionLocal() as db:
        task = models.AnalysisTask(project_id=project_id, material_ids=json.dumps([material['id']]), status='running')
        db.add(task)
        db.commit()
        task_id = task.id
    with TestClient(main.app) as restarted:
        result = restarted.get(f'/api/v1/projects/{project_id}/requirement-tasks/{task_id}').json()
        assert result['status'] == 'interrupted'
        assert restarted.get(f"/api/v1/projects/{project_id}/materials/{material['id']}/original").content == b'example'

def test_question_is_shown_with_source(client, monkeypatch):
    import json
    import app.analysis as analysis
    import app.database as database
    monkeypatch.setattr(analysis, 'SessionLocal', database.SessionLocal)
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'fake-for-test')
    project_id = client.post('/api/v1/projects', json={'name': '问题测试'}).json()['id']
    material_id = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('a.txt', '日期待确认'.encode(), 'text/plain')}).json()['id']
    class Response:
        status_code = 200
        def raise_for_status(self): pass
        def json(self): return {'usage': {'prompt_tokens': 10, 'completion_tokens': 5}, 'choices': [{'message': {'content': json.dumps({'requirements': [{'text': '明确日期', 'material_id': material_id, 'quote': '日期待确认'}], 'questions': [{'text': '截止日期是什么？', 'material_id': material_id, 'quote': '日期待确认'}]})}}]}
    monkeypatch.setattr(analysis.httpx, 'post', lambda *a, **kw: Response())
    task = client.post(f'/api/v1/projects/{project_id}/requirement-tasks', json={'material_ids': [material_id], 'confirm_external_use': True}).json()
    result = client.get(f"/api/v1/projects/{project_id}/requirement-tasks/{task['id']}").json()
    questions = client.get(f'/api/v1/projects/{project_id}/questions').json()
    assert result['usage_complete'] is True and result['estimated_cost_usd'] == 0.000009
    assert questions[0]['text'] == '截止日期是什么？' and questions[0]['quote'] == '日期待确认'
    question_url = f"/api/v1/projects/{project_id}/questions/{questions[0]['id']}"
    assert client.patch(question_url, json={'status': 'resolved', 'answer': ''}).status_code == 422
    assert client.patch(question_url, json={'status': 'answered', 'answer': '客户下周确认'}).status_code == 200
    saved = client.get(f'/api/v1/projects/{project_id}/questions').json()[0]
    assert saved['status'] == 'answered' and saved['answer'] == '客户下周确认'
    requirement = client.get(f'/api/v1/projects/{project_id}/requirements').json()[0]
    assert client.patch(f"/api/v1/projects/{project_id}/requirements/{requirement['id']}", json={'status': 'confirmed', 'text': '人工修改的需求', 'version': 1}).status_code == 200
    history = client.get(f"/api/v1/projects/{project_id}/requirements/{requirement['id']}/history").json()
    assert history[0]['old_text'] == '明确日期' and history[0]['new_text'] == '人工修改的需求'
    client.post(f'/api/v1/projects/{project_id}/baselines')
    baseline = client.get(f'/api/v1/projects/{project_id}/baselines').json()[0]
    assert baseline['questions'][0]['answer'] == '客户下周确认'
    assert baseline['requirements'][0]['text'] == '人工修改的需求'

def test_large_material_is_previewed_without_sending_full_text(client):
    project_id = client.post('/api/v1/projects', json={'name': '大文件'}).json()['id']
    content = '甲' * 50000 + '末尾'
    uploaded = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('source.txt', content.encode(), 'text/plain')}).json()
    assert uploaded['characters'] == len(content)
    assert uploaded['preview_truncated'] is True and len(uploaded['text']) == 10000
    page = client.get(f"/api/v1/projects/{project_id}/materials/{uploaded['id']}/content?offset=50000").json()
    assert page['text'] == '末尾' and page['next_offset'] is None
    invalid = client.get(f"/api/v1/projects/{project_id}/materials/{uploaded['id']}/content?offset=-1")
    assert invalid.status_code == 422

def test_unexpected_errors_do_not_expose_details(client, monkeypatch):
    import app.main as main
    def broken(*args):
        raise RuntimeError('private-document-and-key')
    monkeypatch.setattr(main, 'project_or_404', broken)
    response = client.get('/api/v1/projects/anything/materials')
    assert response.status_code == 500
    assert response.json() == {'error': {'code': 'INTERNAL_ERROR', 'message': '操作未完成，请稍后重试。'}}
    assert 'private-document-and-key' not in response.text

def test_model_needs_explicit_confirmation(client, monkeypatch):
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'fake-for-test')
    project_id = client.post('/api/v1/projects', json={'name': '确认测试'}).json()['id']
    material_id = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('a.txt', b'example', 'text/plain')}).json()['id']
    response = client.post(f'/api/v1/projects/{project_id}/requirement-tasks', json={'material_ids': [material_id]})
    assert response.status_code == 400 and response.json()['error']['code'] == 'CONFIRMATION_REQUIRED'
    assert client.get(f'/api/v1/projects/{project_id}/requirement-tasks').json() == []

def test_grouped_materials_keep_facts_and_all_sources(client, monkeypatch):
    import json
    import app.analysis as analysis
    import app.database as database
    monkeypatch.setattr(analysis, 'SessionLocal', database.SessionLocal)
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'fake-for-test')
    project_id = client.post('/api/v1/projects', json={'name': '多资料测试'}).json()['id']
    first = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('a.txt', '需要项目台账'.encode(), 'text/plain')}).json()['id']
    second = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('b.txt', '共42条账号，其中36条内部账号；也需要项目台账'.encode(), 'text/plain')}).json()['id']
    calls = []
    class Response:
        status_code = 200
        def raise_for_status(self): pass
        def json(self):
            result = {'requirements': [{'text': '项目台账', 'material_id': first, 'quote': '需要项目台账', 'related_sources': [{'material_id': second, 'quote': '也需要项目台账'}]}], 'facts': [{'text': '42条账号中36条为内部账号', 'material_id': second, 'quote': '共42条账号，其中36条内部账号'}], 'questions': []}
            return {'usage': {'prompt_tokens': 20, 'completion_tokens': 10}, 'choices': [{'message': {'content': json.dumps(result)}}]}
    def fake_post(*args, **kwargs):
        assert kwargs['json']['thinking'] == {'type': 'disabled'}
        calls.append(json.loads(kwargs['json']['messages'][1]['content']))
        return Response()
    monkeypatch.setattr(analysis.httpx, 'post', fake_post)
    task = client.post(f'/api/v1/projects/{project_id}/requirement-tasks', json={'material_ids': [first, second], 'confirm_external_use': True}).json()
    result = client.get(f"/api/v1/projects/{project_id}/requirement-tasks/{task['id']}").json()
    assert result['status'] == 'completed' and result['segments_completed'] == 2 and result['calls'] == 1
    assert set(calls[0]) == {'S1', 'S2'}
    assert set(calls[0].values()) == {'需要项目台账', '共42条账号，其中36条内部账号；也需要项目台账'}
    requirement = client.get(f'/api/v1/projects/{project_id}/requirements').json()[0]
    assert requirement['related_sources'] == [{'material_id': second, 'quote': '也需要项目台账'}]
    facts = client.get(f'/api/v1/projects/{project_id}/facts').json()
    assert facts[0]['text'] == '42条账号中36条为内部账号'

def test_invalid_additional_citation_rejects_entire_group(client, monkeypatch):
    import json
    import app.analysis as analysis
    import app.database as database
    monkeypatch.setattr(analysis, 'SessionLocal', database.SessionLocal)
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'fake-for-test')
    project_id = client.post('/api/v1/projects', json={'name': '出处检查'}).json()['id']
    material_id = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('a.txt', b'known source', 'text/plain')}).json()['id']
    class Response:
        status_code = 200
        def raise_for_status(self): pass
        def json(self):
            result = {'requirements': [{'text': 'test', 'material_id': material_id, 'quote': 'known', 'related_sources': [{'material_id': material_id, 'quote': 'fabricated'}]}], 'facts': [], 'questions': []}
            return {'choices': [{'message': {'content': json.dumps(result)}}]}
    monkeypatch.setattr(analysis.httpx, 'post', lambda *a, **kw: Response())
    task = client.post(f'/api/v1/projects/{project_id}/requirement-tasks', json={'material_ids': [material_id], 'confirm_external_use': True}).json()
    assert client.get(f"/api/v1/projects/{project_id}/requirement-tasks/{task['id']}").json()['status'] == 'failed'
    assert client.get(f'/api/v1/projects/{project_id}/requirements').json() == []

def test_truncated_output_is_not_saved_as_complete(client, monkeypatch):
    import app.analysis as analysis
    import app.database as database
    monkeypatch.setattr(analysis, 'SessionLocal', database.SessionLocal)
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'fake-for-test')
    project_id = client.post('/api/v1/projects', json={'name': '截断测试'}).json()['id']
    material_id = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('a.txt', b'sample', 'text/plain')}).json()['id']
    class Response:
        status_code = 200
        def raise_for_status(self): pass
        def json(self): return {'usage': {'prompt_tokens': 10, 'completion_tokens': 16384}, 'choices': [{'finish_reason': 'length', 'message': {'content': '{'}}]}
    monkeypatch.setattr(analysis.httpx, 'post', lambda *a, **kw: Response())
    task = client.post(f'/api/v1/projects/{project_id}/requirement-tasks', json={'material_ids': [material_id], 'confirm_external_use': True}).json()
    result = client.get(f"/api/v1/projects/{project_id}/requirement-tasks/{task['id']}").json()
    assert result['status'] == 'failed' and result['segments_completed'] == 0
    assert result['calls'] == 1 and '截断' in result['error']
    assert client.get(f'/api/v1/projects/{project_id}/requirements').json() == []

def test_short_source_aliases_resolve_to_the_correct_documents(client, monkeypatch):
    import json
    import app.analysis as analysis
    import app.database as database
    monkeypatch.setattr(analysis, 'SessionLocal', database.SessionLocal)
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'fake-for-test')
    project_id = client.post('/api/v1/projects', json={'name': '编号校验'}).json()['id']
    first = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('a.txt', b'alpha unique', 'text/plain')}).json()['id']
    second = client.post(f'/api/v1/projects/{project_id}/materials', files={'file': ('b.txt', b'beta unique', 'text/plain')}).json()['id']
    class Response:
        status_code = 200
        def __init__(self, sources): self.sources = sources
        def raise_for_status(self): pass
        def json(self):
            alpha = next(k for k, v in self.sources.items() if 'alpha' in v)
            beta = next(k for k, v in self.sources.items() if 'beta' in v)
            result = {'requirements': [{'text': 'combined', 'material_id': alpha, 'quote': 'alpha unique', 'related_sources': [{'material_id': beta, 'quote': 'beta unique'}]}], 'facts': [], 'questions': []}
            return {'choices': [{'message': {'content': json.dumps(result)}}]}
    def fake_post(*args, **kwargs):
        return Response(json.loads(kwargs['json']['messages'][1]['content']))
    monkeypatch.setattr(analysis.httpx, 'post', fake_post)
    task = client.post(f'/api/v1/projects/{project_id}/requirement-tasks', json={'material_ids': [first, second], 'confirm_external_use': True}).json()
    result = client.get(f"/api/v1/projects/{project_id}/requirement-tasks/{task['id']}").json()
    assert result['status'] == 'completed'
    requirement = client.get(f'/api/v1/projects/{project_id}/requirements').json()[0]
    assert requirement['material_id'] == first
    assert requirement['related_sources'] == [{'material_id': second, 'quote': 'beta unique'}]

def test_fact_review_can_be_saved_reverted_and_restored(client):
    import json
    from sqlalchemy import select
    import app.database as database
    import app.models as models
    import app.main as main
    project = client.post('/api/v1/projects', json={'name': '核对测试'}).json()['id']
    other = client.post('/api/v1/projects', json={'name': '另一个项目'}).json()['id']
    material = client.post(f'/api/v1/projects/{project}/materials', files={'file': ('a.txt', '原始事实'.encode(), 'text/plain')}).json()['id']
    with database.SessionLocal() as db:
        task = models.AnalysisTask(project_id=project, status='completed', material_ids=json.dumps([material]))
        db.add(task); db.flush()
        fact = models.ExtractedFact(project_id=project, task_id=task.id, material_id=material, text='事实', quote='原始事实')
        db.add(fact); db.commit(); fact_id = fact.id
    path = f'/api/v1/projects/{project}/facts/{fact_id}/review'
    before = client.get(f'/api/v1/projects/{project}/facts').json()[0]
    assert before['review_status'] == 'pending' and before['review_version'] == 0
    assert client.patch(f'/api/v1/projects/{other}/facts/{fact_id}/review', json={'review_status': 'reviewed', 'review_version': 0}).status_code == 404
    marked = client.patch(path, json={'review_status': 'reviewed', 'review_version': 0}).json()
    assert marked['review_status'] == 'reviewed' and marked['review_version'] == 1
    assert client.patch(path, json={'review_status': 'pending', 'review_version': 0}).status_code == 409
    with TestClient(main.app) as restarted:
        saved = restarted.get(f'/api/v1/projects/{project}/facts').json()[0]
        assert saved['review_status'] == 'reviewed' and saved['review_version'] == 1
        reverted = restarted.patch(path, json={'review_status': 'pending', 'review_version': 1}).json()
        assert reverted['review_status'] == 'pending' and reverted['review_version'] == 2
        assert restarted.get(f'/api/v1/projects/{project}/baselines').json() == []
    with database.SessionLocal() as db:
        events = list(db.scalars(select(models.FactReviewEvent).where(models.FactReviewEvent.fact_id == fact_id)))
        assert [(event.old_status, event.new_status) for event in events] == [('pending', 'reviewed'), ('reviewed', 'pending')]

def test_proposal_generation_and_edit_keeps_original(client, monkeypatch):
    import json
    import app.database as database
    import app.models as models
    import app.main as main
    import app.proposal as proposal
    monkeypatch.setattr(proposal, 'SessionLocal', database.SessionLocal)
    monkeypatch.setattr(proposal, 'ProposalDraft', models.ProposalDraft)
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'fake-for-test')
    project = client.post('/api/v1/projects', json={'name':'虚构项目'}).json()['id']
    other = client.post('/api/v1/projects', json={'name':'其他项目'}).json()['id']
    with database.SessionLocal() as db:
        b = models.Baseline(project_id=project, snapshot=json.dumps([{'text':'建立台账'}], ensure_ascii=False))
        db.add(b); db.commit(); baseline_id = b.id
    class Response:
        status_code = 200
        def raise_for_status(self): pass
        def json(self): return {'usage': {'prompt_tokens':20,'completion_tokens':10}, 'choices':[{'finish_reason':'stop','message':{'content':json.dumps({'content':'# 方案初稿\n\n建立台账。'},ensure_ascii=False)}}]}
    monkeypatch.setattr(proposal.httpx, 'post', lambda *a, **k: Response())
    url = f'/api/v1/projects/{project}/proposals'
    body = {'baseline_id':baseline_id,'template_id':'demo-v1','model_name':'deepseek-flash','include_example':True,'confirm_external_use':True}
    assert client.post(url, json={**body,'confirm_external_use':False}).status_code == 400
    assert client.post(f'/api/v1/projects/{other}/proposals', json=body).status_code == 404
    started = client.post(url, json=body)
    assert started.status_code == 202
    draft = client.get(url).json()[0]
    assert draft['status'] == 'completed' and draft['calls'] == 1 and draft['prompt_tokens'] == 20
    assert '建立台账' in draft['content']
    edit_url = url + '/' + draft['id']
    edited = client.patch(edit_url, json={'content':'# 人工修改','version':draft['version']}).json()
    assert edited['content'] == '# 人工修改' and edited['version'] == 2
    assert client.patch(edit_url, json={'content':'旧版本覆盖','version':1}).status_code == 409
    with database.SessionLocal() as db:
        from sqlalchemy import select
        previous = db.scalar(select(models.ProposalRevision).where(models.ProposalRevision.draft_id == draft['id']))
        assert previous.version == 1 and '建立台账' in previous.content
        saved_draft = db.get(models.ProposalDraft, draft['id'])
        assert '虚构参考案例' in saved_draft.input_snapshot
    with TestClient(main.app) as restarted:
        assert restarted.get(url).json()[0]['content'] == '# 人工修改'

def test_proposal_rejects_invalid_model_output(client, monkeypatch):
    import json
    import app.database as database
    import app.models as models
    import app.proposal as proposal
    monkeypatch.setattr(proposal, 'SessionLocal', database.SessionLocal)
    monkeypatch.setattr(proposal, 'ProposalDraft', models.ProposalDraft)
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'fake-for-test')
    project = client.post('/api/v1/projects', json={'name':'虚构项目'}).json()['id']
    with database.SessionLocal() as db:
        b = models.Baseline(project_id=project, snapshot=json.dumps([{'text':'建立台账'}]))
        db.add(b); db.commit(); baseline_id = b.id
    class Response:
        status_code = 200
        def raise_for_status(self): pass
        def json(self): return {'choices':[{'finish_reason':'stop','message':{'content':'{"wrong":1}'}}]}
    monkeypatch.setattr(proposal.httpx, 'post', lambda *a, **k: Response())
    url = f'/api/v1/projects/{project}/proposals'
    client.post(url, json={'baseline_id':baseline_id,'template_id':'demo-v1','confirm_external_use':True})
    draft = client.get(url).json()[0]
    assert draft['status'] == 'failed' and draft['content'] == ''

def test_proposal_input_uses_filenames_instead_of_internal_ids(client, monkeypatch):
    import json
    import app.database as database
    import app.models as models
    import app.proposal as proposal
    monkeypatch.setattr(proposal, 'SessionLocal', database.SessionLocal)
    monkeypatch.setattr(proposal, 'ProposalDraft', models.ProposalDraft)
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'fake-for-test')
    project = client.post('/api/v1/projects', json={'name':'虚构项目'}).json()['id']
    material = client.post(f'/api/v1/projects/{project}/materials', files={'file': ('访谈.txt', '需要统一台账。'.encode(), 'text/plain')}).json()
    missing_id = '11111111-1111-4111-8111-111111111111'
    with database.SessionLocal() as db:
        entry = {'id':'内部需求编号', 'text':'建立台账', 'material_id':material['id'], 'quote':'需要统一台账。', 'related_sources':[{'material_id':missing_id,'quote':'待核对原文'}]}
        b = models.Baseline(project_id=project, snapshot=json.dumps([entry],ensure_ascii=False))
        db.add(b); db.commit(); baseline_id = b.id
        db.add(models.BaselineQuestions(baseline_id=baseline_id, snapshot=json.dumps([{'id':'内部问题编号','text':'具体规则？','status':'pending','material_id':material['id'],'quote':'需要统一台账。','answer':'','related_sources':[]}],ensure_ascii=False))); db.commit()
    captured = []
    class Response:
        status_code = 200
        def raise_for_status(self): pass
        def json(self): return {'choices':[{'finish_reason':'stop','message':{'content':'{"content":"方案初稿"}'}}]}
    def fake_post(*args, **kwargs):
        captured.append(kwargs['json']['messages'][1]['content'])
        return Response()
    monkeypatch.setattr(proposal.httpx, 'post', fake_post)
    result = client.post(f'/api/v1/projects/{project}/proposals', json={'baseline_id':baseline_id,'template_id':'demo-v1','confirm_external_use':True})
    assert result.status_code == 202 and len(captured) == 1
    sent = captured[0]
    assert '访谈.txt' in sent and '需要统一台账。' in sent and '来源资料待核对' in sent
    assert material['id'] not in sent and missing_id not in sent and '内部需求编号' not in sent and '内部问题编号' not in sent

def make_review_fixture(client):
    import json
    import app.database as database
    import app.models as models
    import app.main as main
    project = client.post('/api/v1/projects', json={'name':'虚构项目'}).json()['id']
    other = client.post('/api/v1/projects', json={'name':'其他项目'}).json()['id']
    material = client.post(f'/api/v1/projects/{project}/materials', files={'file':('访谈.txt', '客户希望建立台账。'.encode(), 'text/plain')}).json()
    with database.SessionLocal() as db:
        task=models.AnalysisTask(project_id=project,status='completed',material_ids=json.dumps([material['id']]));db.add(task);db.flush()
        req=models.Requirement(project_id=project,task_id=task.id,material_id=material['id'],text='建立台账',original_text='建立台账',quote='客户希望建立台账。',status='confirmed');db.add(req)
        q=models.OpenQuestion(id='问题一',project_id=project,task_id=task.id,material_id=material['id'],text='具体实施规则？',quote='客户希望建立台账。',status='pending');db.add(q);db.flush()
        requirement = {'id':req.id,'version':req.version,'text':'建立台账','material_id':material['id'],'quote':'客户希望建立台账。','related_sources':[]}
        question = {'id':'问题一','text':'具体实施规则？','status':'pending','answer':'','material_id':material['id'],'quote':'客户希望建立台账。','related_sources':[]}
        baseline = models.Baseline(project_id=project,snapshot=json.dumps([requirement],ensure_ascii=False));db.add(baseline);db.flush()
        db.add(models.BaselineQuestions(baseline_id=baseline.id,snapshot=json.dumps([question],ensure_ascii=False)))
        content = '# 虚构方案\n\n## 功能\n\n| 问题 | 需补充资料 |\n|---|---|\n| 实施规则 | '+material['id']+' |\n\n正文。'
        draft = models.ProposalDraft(project_id=project,baseline_id=baseline.id,template_id='demo-v1',template_snapshot=main.TEMPLATE_PATH.read_text(encoding='utf-8'),input_snapshot='{}',model_name='deepseek-flash',status='completed',content=content,calls=0)
        db.add(draft);db.commit(); draft_id=draft.id
    return project, other, material['id'], draft_id

def test_review_export_gates_and_version_restore(client):
    import io
    import zipfile
    from docx import Document
    import app.database as database
    import app.models as models
    import app.main as main
    from sqlalchemy import select
    project, other, material, draft_id = make_review_fixture(client)
    base=f'/api/v1/projects/{project}/proposals/{draft_id}'
    assert client.get(f'/api/v1/projects/{other}/proposals/{draft_id}/review').status_code == 404
    assert client.post(base+'/export',json={'version':1,'kind':'approved'}).status_code == 409
    response=client.post(base+'/export',json={'version':1,'kind':'draft'})
    assert response.status_code == 200
    doc=Document(io.BytesIO(response.content)); text='\n'.join(p.text for p in doc.paragraphs)
    cells=' '.join(c.text for table in doc.tables for row in table.rows for c in row.cells)
    assert '草稿（未审核）' in text and '未解决事项' in text and '访谈.txt' in text
    assert '已有资料' in cells and '还需补充' in cells and '访谈.txt' in cells and material not in cells
    with zipfile.ZipFile(io.BytesIO(response.content)) as zipped:
        assert '目录' in text and b'tblBorders' in zipped.read('word/document.xml') and b'PAGE' in zipped.read('word/footer1.xml')
    assert client.post(base+'/review',json={'version':1,'action':'submit'}).json()['status']=='submitted'
    payload={'version':1,'action':'approve','note':'已核对','other_critical_checked':True,'decisions':[]}
    assert client.post(base+'/review',json=payload).status_code==409
    payload['decisions']=[{'question_id':'问题一','classification':'critical','disposition':'general_approved','reason':'需要补充'}]
    assert client.post(base+'/review',json=payload).status_code==409
    payload['decisions'][0]['classification']='not_classified'
    assert client.post(base+'/review',json=payload).status_code==409
    payload['decisions'][0]={'question_id':'问题一','classification':'general','disposition':'general_approved','reason':'不影响主要结论，保留待确认'}
    payload['decisions'].append({'question_id':'extra-test','issue_text':'方案新增的接口问题','classification':'critical','disposition':'general_approved','reason':'待核实'})
    assert client.post(base+'/review',json=payload).status_code==409
    payload['decisions'][1]['disposition']='noncommit'
    payload['decisions'][1]['classification']='not_classified'
    assert client.post(base+'/review',json=payload).json()['status']=='approved'
    history=client.get(base+'/review-history').json()
    assert [entry['action'] for entry in history] == ['approved','submitted'] and history[0]['decisions'][0]['reason']
    approved=client.post(base+'/export',json={'version':1,'kind':'approved'})
    assert approved.status_code==200
    approved_text='\n'.join(p.text for p in Document(io.BytesIO(approved.content)).paragraphs)
    assert '已审核方案（本机试用）' in approved_text and '一般事项批准保留（未解决）' in approved_text and '方案新增的接口问题' in approved_text
    with database.SessionLocal() as db:
        assert len(list(db.scalars(select(models.ProposalExportEvent)))) == 2
        source=db.get(models.Material,material); original=source.text; source.text += '资料发生变化。'; db.commit()
    assert client.get(base+'/review').json()['status']=='needs_review'
    assert client.post(base+'/export',json={'version':1,'kind':'approved'}).status_code==409
    with database.SessionLocal() as db:
        source=db.get(models.Material,material); source.text=original; db.commit()
    current=client.get(f'/api/v1/projects/{project}/requirements').json()[0]
    assert client.patch(f'/api/v1/projects/{project}/requirements/{current["id"]}',json={'text':'更新后的需求','status':'confirmed','version':current['version']}).status_code==200
    assert client.get(base+'/review').json()['status']=='needs_review'
    assert client.post(base+'/export',json={'version':1,'kind':'approved'}).status_code==409
    assert client.patch(base,json={'content':'# 修改后的方案','version':1}).json()['version']==2
    assert client.get(base+'/review').json()['status']=='draft'
    assert client.post(base+'/export',json={'version':2,'kind':'approved'}).status_code==409
    with TestClient(main.app) as restarted:
        assert restarted.get(base+'/review').json()['status']=='draft'
        with database.SessionLocal() as db:
            assert len(list(db.scalars(select(models.ProposalReviewEvent))))==2

def test_review_blocks_changed_or_missing_source(client):
    import app.database as database
    import app.models as models
    project, other, material, draft_id = make_review_fixture(client)
    base=f'/api/v1/projects/{project}/proposals/{draft_id}'
    assert client.post(base+'/review',json={'version':1,'action':'submit'}).status_code==200
    with database.SessionLocal() as db:
        row=db.get(models.Material,material);row.text += '资料发生变化。';db.commit()
    state=client.get(base+'/review').json()
    assert state['status']=='needs_review'
    payload={'version':1,'action':'approve','note':'已核对','other_critical_checked':True,'decisions':[{'question_id':'问题一','classification':'critical','disposition':'resolved','reason':'已处理'}]}
    assert client.post(base+'/review',json=payload).status_code==409
    assert client.post(base+'/export',json={'version':1,'kind':'approved'}).status_code==409
    returned=client.post(base+'/review',json={'version':1,'action':'return','note':'资料变了，请重新核对。'})
    assert returned.status_code==200 and returned.json()['status']=='returned'

def test_review_blocks_changed_question(client):
    project, _, _, draft_id = make_review_fixture(client)
    base=f'/api/v1/projects/{project}/proposals/{draft_id}'
    assert client.post(base+'/review',json={'version':1,'action':'submit'}).status_code==200
    assert client.patch(f'/api/v1/projects/{project}/questions/问题一',json={'status':'answered','answer':'客户正在核对'}).status_code==200
    assert client.get(base+'/review').json()['status']=='needs_review'
    payload={'version':1,'action':'approve','note':'已核对','other_critical_checked':True,'decisions':[{'question_id':'问题一','classification':'critical','disposition':'resolved','reason':'已处理'}]}
    assert client.post(base+'/review',json=payload).status_code==409

def test_old_baseline_without_question_snapshot_cannot_be_approved(client):
    import io
    from docx import Document
    import app.database as database
    import app.models as models
    project, _, _, draft_id = make_review_fixture(client)
    base=f'/api/v1/projects/{project}/proposals/{draft_id}'
    with database.SessionLocal() as db:
        draft=db.get(models.ProposalDraft,draft_id); record=db.get(models.BaselineQuestions,draft.baseline_id); db.delete(record); db.commit()
    assert client.post(base+'/review',json={'version':1,'action':'submit'}).status_code==409
    response=client.post(base+'/export',json={'version':1,'kind':'draft'})
    assert response.status_code==200
    assert '没有待确认问题记录' in '\n'.join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)

def test_missing_original_is_not_exported_as_verified_quote(client):
    import io
    from docx import Document
    import app.database as database
    project, _, material, draft_id = make_review_fixture(client)
    base=f'/api/v1/projects/{project}/proposals/{draft_id}'
    (database.DATA_DIR / 'materials' / material).unlink()
    assert client.post(base+'/review',json={'version':1,'action':'submit'}).status_code==409
    response=client.post(base+'/export',json={'version':1,'kind':'draft'})
    assert response.status_code==200
    text='\n'.join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)
    assert '资料不可用或已变化，请核对' in text
    assert '原文摘录：客户希望建立台账。' not in text


def test_reviewed_fact_snapshot_is_used_with_sources_and_revocation_blocks_approval(client, monkeypatch):
    import json
    import app.database as database
    import app.models as models
    import app.proposal as proposal
    monkeypatch.setattr(proposal, 'SessionLocal', database.SessionLocal)
    monkeypatch.setattr(proposal, 'ProposalDraft', models.ProposalDraft)
    monkeypatch.setenv('DEEPSEEK_API_KEY', 'fake-for-test')
    project = client.post('/api/v1/projects', json={'name':'事实快照测试'}).json()['id']
    text = '需要统一台账。五个科室使用Excel。另有未核对资料。'
    material = client.post(f'/api/v1/projects/{project}/materials', files={'file':('访谈.txt',text.encode(),'text/plain')}).json()['id']
    with database.SessionLocal() as db:
        task = models.AnalysisTask(project_id=project, status='completed', material_ids=json.dumps([material])); db.add(task); db.flush()
        req = models.Requirement(project_id=project, task_id=task.id, material_id=material, text='统一台账', original_text='统一台账', quote='需要统一台账。', status='confirmed'); db.add(req)
        a = models.ExtractedFact(project_id=project, task_id=task.id, material_id=material, text='五个科室使用Excel', quote='五个科室使用Excel。')
        b = models.ExtractedFact(project_id=project, task_id=task.id, material_id=material, text='未核对的事实', quote='另有未核对资料。'); db.add_all([a,b]); db.commit(); a_id, b_id = a.id, b.id
    path = f'/api/v1/projects/{project}/facts/{a_id}/review'
    assert client.patch(path,json={'review_status':'reviewed','review_version':0}).status_code == 200
    baseline_id = client.post(f'/api/v1/projects/{project}/baselines').json()['id']
    snapshot = client.get(f'/api/v1/projects/{project}/baselines').json()[0]
    assert snapshot['facts_snapshot_available'] and [f['id'] for f in snapshot['facts']] == [a_id]
    assert client.patch(path,json={'review_status':'pending','review_version':1}).status_code == 200
    assert client.patch(f'/api/v1/projects/{project}/facts/{b_id}/review',json={'review_status':'reviewed','review_version':0}).status_code == 200
    assert [f['id'] for f in client.get(f'/api/v1/projects/{project}/baselines').json()[0]['facts']] == [a_id]
    sent = []
    class Response:
        status_code = 200
        def raise_for_status(self): pass
        def json(self): return {'choices':[{'finish_reason':'stop','message':{'content':'{"content":"方案初稿"}'}}]}
    def fake_post(*args, **kwargs):
        sent.append(json.loads(kwargs['json']['messages'][1]['content'])); return Response()
    monkeypatch.setattr(proposal.httpx, 'post', fake_post)
    r=client.post(f'/api/v1/projects/{project}/proposals',json={'baseline_id':baseline_id,'template_id':'demo-v1','confirm_external_use':True})
    assert r.status_code == 202
    draft_id=r.json()['id']
    assert sent[0]['项目和需求']['已核对资料事实'] == [{'事实内容':'五个科室使用Excel','已有资料出处':[{'资料名称':'访谈.txt'}]}]
    assert '未核对的事实' not in json.dumps(sent[0],ensure_ascii=False)
    assert a_id not in json.dumps(sent[0],ensure_ascii=False)
    submit=client.post(f'/api/v1/projects/{project}/proposals/{draft_id}/review',json={'version':1,'action':'submit'})
    assert submit.status_code == 409 and '资料事实' in submit.json()['error']['message']
    new_id=client.post(f'/api/v1/projects/{project}/baselines').json()['id']
    new=next(x for x in client.get(f'/api/v1/projects/{project}/baselines').json() if x['id']==new_id)
    assert [f['id'] for f in new['facts']] == [b_id]
