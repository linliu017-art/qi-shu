from contextlib import contextmanager
import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from app.database import Base
import app.main as main
from app.models import ProductWorkRevision, Project, UserAccount
from app.membership import initialize_project_owner, add_member, remove_member
from app.security_context import actor_id, multiuser

@contextmanager
def identity(user):
    a = actor_id.set(user); m = multiuser.set(True)
    try:
        yield
    finally:
        multiuser.reset(m); actor_id.reset(a)

def test_personal_and_project_access_and_versioning():
    engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine, expire_on_commit=False) as db:
        owner = UserAccount(username='owner', password_hash='x', active=True)
        other = UserAccount(username='other', password_hash='x', active=True)
        db.add_all([owner, other]); db.flush()
        project = Project(name='测试', background='')
        initialize_project_owner(db, project, owner.id); db.commit()
        with identity(owner.id):
            work = main.create_product_work(main.ProductWorkCreate(title=' 竞品研究 '), db)
            first = main.edit_product_work(work.id, main.ProductWorkEdit(title='竞品研究', content='观察', version=1), db)
            assert first.version == 2
            with pytest.raises(main.ApiError) as stale:
                main.edit_product_work(work.id, main.ProductWorkEdit(title='冲突', content='不能覆盖', version=1), db)
            assert stale.value.status == 409
            assert main.get_product_work(work.id, db).content == '观察'
            assert db.scalar(select(ProductWorkRevision).where(ProductWorkRevision.work_id == work.id)).version == 1
        with identity(other.id):
            assert main.list_product_work(db) == []
            with pytest.raises(main.ApiError) as hidden:
                main.get_product_work(work.id, db)
            assert hidden.value.status == 404
        with identity(owner.id):
            with pytest.raises(main.ApiError) as no_consent:
                main.assign_product_work(work.id, main.ProductWorkMove(project_id=project.id, version=2), db)
            assert no_consent.value.code == 'CONFIRMATION_REQUIRED'
            assigned = main.assign_product_work(work.id, main.ProductWorkMove(project_id=project.id, version=2, confirm_visibility=True), db)
            assert assigned.project_id == project.id and assigned.version == 3
        with identity(other.id):
            with pytest.raises(main.ApiError) as no_project:
                main.get_product_work(work.id, db)
            assert no_project.value.status == 404
        add_member(db, project.id, owner.id, other.id)
        with identity(owner.id):
            initiative=main.create_product_work(main.ProductWorkCreate(title='立项',project_id=project.id,kind='initiation'),db)
        with identity(other.id):
            assert main.get_product_work(work.id, db).content == '观察'
            assert main.get_product_work(work.id, db).editable is False
            with pytest.raises(main.ApiError) as forbidden:
                main.edit_product_work(work.id, main.ProductWorkEdit(title='改写', content='错误', version=3), db)
            assert forbidden.value.status == 403
            assert main.get_product_work(initiative.id,db).kind=='initiation'
            with pytest.raises(main.ApiError) as cannot_change:
                main.set_initiation_status(initiative.id,main.InitiationStatusInput(status='approved',version=1),db)
            assert cannot_change.value.status == 403
        remove_member(db, project.id, owner.id, other.id)
        with identity(other.id):
            with pytest.raises(main.ApiError) as removed:
                main.get_product_work(work.id, db)
            assert removed.value.status == 404
    engine.dispose()

def test_http_create_save_history_and_export():
    from fastapi.testclient import TestClient
    from app.database import Base as CurrentBase
    engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
    CurrentBase.metadata.create_all(engine)
    def test_db():
        with Session(engine, expire_on_commit=False) as db:
            yield db
    main.app.dependency_overrides[main.db_session] = test_db
    try:
        client = TestClient(main.app)
        result = client.post('/api/v1/product-work', json={'title':'测试竞品'})
        assert result.status_code == 201
        work = result.json()
        result = client.put('/api/v1/product-work/'+work['id'], json={'title':'测试竞品','content':'# 观察\n公开资料','version':1})
        assert result.status_code == 200 and result.json()['version'] == 2
        assert client.get('/api/v1/product-work/'+work['id']+'/revisions').json()[0]['version'] == 1
        exported = client.get('/api/v1/product-work/'+work['id']+'/export')
        assert exported.status_code == 200 and exported.content.decode('utf-8') == '# 观察\n公开资料'
    finally:
        main.app.dependency_overrides.pop(main.db_session, None)
        engine.dispose()

def test_research_requires_confirmation_and_returns_candidate(monkeypatch):
    from fastapi.testclient import TestClient
    from app.database import Base as CurrentBase
    import app.product_research as research
    engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
    CurrentBase.metadata.create_all(engine)
    def session():
        return Session(engine, expire_on_commit=False)
    def test_db():
        with session() as db: yield db
    monkeypatch.setattr(main, 'model_key', lambda: 'fake')
    monkeypatch.setattr(main, 'model_settings', lambda db: (['deepseek-test'], 1))
    monkeypatch.setattr(research, 'model_key', lambda: 'fake')
    monkeypatch.setattr(research, 'model_settings', lambda db: (['deepseek-test'], 1))
    monkeypatch.setattr(research, 'SessionLocal', session)
    class Reply:
        def raise_for_status(self): pass
        def json(self): return {'choices':[{'finish_reason':'stop','message':{'content':'{"report":"测试报告（未实测）"}'}}]}
    monkeypatch.setattr(research.httpx, 'post', lambda *args, **kwargs: Reply())
    main.app.dependency_overrides[main.db_session] = test_db
    try:
        client = TestClient(main.app)
        work = client.post('/api/v1/product-work', json={'title':'虚构竞品'}).json()
        payload={'goal':'分析定位','source_text':'虚构资料','version':1,'model':'deepseek-test'}
        denied=client.post('/api/v1/product-work/'+work['id']+'/research',json=payload)
        assert denied.status_code == 400 and denied.json()['error']['code']=='CONFIRMATION_REQUIRED'
        payload['confirm_external_use']=True
        started=client.post('/api/v1/product-work/'+work['id']+'/research',json=payload)
        assert started.status_code == 202
        tasks=client.get('/api/v1/product-work/'+work['id']+'/research').json()
        assert tasks[0]['status']=='completed' and tasks[0]['result']=='测试报告（未实测）'
        assert client.get('/api/v1/product-work/'+work['id']).json()['content']==''
        def cancel_during_call(*args, **kwargs):
            with session() as db:
                active=db.scalar(select(main.ProductResearchTask).where(main.ProductResearchTask.status=='running'))
                active.status='cancelled'; db.commit()
            return Reply()
        monkeypatch.setattr(research.httpx,'post',cancel_during_call)
        started=client.post('/api/v1/product-work/'+work['id']+'/research',json=payload)
        assert started.status_code==202
        cancelled=next(t for t in client.get('/api/v1/product-work/'+work['id']+'/research').json() if t['id']==started.json()['id'])
        assert cancelled['status']=='cancelled' and cancelled['result']==''
    finally:
        main.app.dependency_overrides.pop(main.db_session,None)
        engine.dispose()

def test_initiation_is_independent_and_status_is_manual():
    from fastapi.testclient import TestClient
    from app.database import Base as CurrentBase
    engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
    CurrentBase.metadata.create_all(engine)
    def test_db():
        with Session(engine, expire_on_commit=False) as db: yield db
    main.app.dependency_overrides[main.db_session]=test_db
    try:
        client=TestClient(main.app)
        created=client.post('/api/v1/product-work',json={'title':'虚构产品立项','kind':'initiation'})
        assert created.status_code==201
        row=created.json(); assert row['kind']=='initiation' and row['status']=='draft'
        saved=client.put('/api/v1/product-work/'+row['id'],json={'title':'虚构产品立项','content':'背景待补充','version':1})
        assert saved.status_code==200 and saved.json()['version']==2
        status=client.post('/api/v1/product-work/'+row['id']+'/status',json={'status':'approved','version':2})
        assert status.status_code==200 and status.json()['status']=='approved' and status.json()['version']==3
        stale=client.post('/api/v1/product-work/'+row['id']+'/status',json={'status':'deferred','version':2})
        assert stale.status_code==409
        history=client.get('/api/v1/product-work/'+row['id']+'/status-history').json()
        assert len(history)==1 and history[0]['old_status']=='draft' and history[0]['new_status']=='approved'
        assert client.get('/api/v1/product-work/'+row['id']).json()['content']=='背景待补充'
        assert client.get('/api/v1/product-work/'+row['id']+'/export').status_code==400
    finally:
        main.app.dependency_overrides.pop(main.db_session,None)
        engine.dispose()

def test_initiation_links_are_project_scoped_and_removable():
    from fastapi.testclient import TestClient
    from app.database import Base as CurrentBase
    engine=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool)
    CurrentBase.metadata.create_all(engine)
    def test_db():
        with Session(engine,expire_on_commit=False) as db: yield db
    main.app.dependency_overrides[main.db_session]=test_db
    try:
        client=TestClient(main.app)
        p=client.post('/api/v1/projects',json={'name':'项目A'}).json()['id']
        other=client.post('/api/v1/projects',json={'name':'项目B'}).json()['id']
        report=client.post('/api/v1/product-work',json={'title':'立项','kind':'initiation','project_id':p}).json()
        source=client.post('/api/v1/product-work',json={'title':'竞品A','project_id':p}).json()
        outside=client.post('/api/v1/product-work',json={'title':'竞品B','project_id':other}).json()
        path='/api/v1/product-work/'+report['id']+'/links'
        assert client.post(path,json={'target_type':'competitor','target_id':outside['id'],'version':1}).status_code==404
        linked=client.post(path,json={'target_type':'competitor','target_id':source['id'],'version':1})
        assert linked.status_code==200 and linked.json()['version']==2
        items=client.get(path).json(); assert len(items)==1 and items[0]['state']=='available'
        updated=client.put('/api/v1/product-work/'+source['id'],json={'title':'竞品A新版','content':'新内容','version':1})
        assert updated.status_code==200
        assert client.get(path).json()[0]['state']=='changed'
        removed=client.post(path+'/'+items[0]['id']+'/remove',json={'version':2})
        assert removed.status_code==200 and client.get(path).json()==[]
        assert client.get('/api/v1/product-work/'+source['id']).json()['content']=='新内容'
    finally:
        main.app.dependency_overrides.pop(main.db_session,None)
        engine.dispose()

def test_research_only_sends_selected_links_and_hides_stale_evidence(monkeypatch):
    from fastapi.testclient import TestClient
    from app.database import Base as CurrentBase
    import app.product_research as research
    engine=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool)
    CurrentBase.metadata.create_all(engine)
    def session():return Session(engine,expire_on_commit=False)
    def test_db():
        with session() as db: yield db
    monkeypatch.setattr(main,'model_key',lambda:'fake')
    monkeypatch.setattr(main,'model_settings',lambda db:(['deepseek-test'],1))
    monkeypatch.setattr(research,'model_key',lambda:'fake')
    monkeypatch.setattr(research,'model_settings',lambda db:(['deepseek-test'],1))
    monkeypatch.setattr(research,'SessionLocal',session)
    captured=[]
    class Reply:
        def raise_for_status(self):pass
        def json(self):return {'choices':[{'finish_reason':'stop','message':{'content':'{"report":"分析候选"}'}}]}
    def fake_post(*args,**kwargs):captured.append(kwargs['json']);return Reply()
    monkeypatch.setattr(research.httpx,'post',fake_post)
    main.app.dependency_overrides[main.db_session]=test_db
    try:
        client=TestClient(main.app)
        project=client.post('/api/v1/projects',json={'name':'虚构项目'}).json()['id']
        report=client.post('/api/v1/product-work',json={'title':'立项','kind':'initiation','project_id':project}).json()
        source=client.post('/api/v1/product-work',json={'title':'指定依据','project_id':project}).json()
        secret=client.post('/api/v1/product-work',json={'title':'未选依据','project_id':project}).json()
        client.put('/api/v1/product-work/'+source['id'],json={'title':'指定依据','content':'允许发送的资料','version':1})
        client.put('/api/v1/product-work/'+secret['id'],json={'title':'未选依据','content':'不可发送的资料','version':1})
        path='/api/v1/product-work/'+report['id']
        assert client.post(path+'/links',json={'target_type':'competitor','target_id':source['id'],'version':1}).status_code==200
        link=client.get(path+'/links').json()[0]
        result=client.post(path+'/research',json={'goal':'起草立项','source_text':'','link_ids':[link['id']],
            'version':2,'model':'deepseek-test','confirm_external_use':True})
        assert result.status_code==202
        sent=str(captured[0]['messages'])
        assert '允许发送的资料' in sent and '不可发送的资料' not in sent
        assert client.get(path+'/research').json()[0]['result']=='分析候选'
        client.put('/api/v1/product-work/'+source['id'],json={'title':'指定依据','content':'依据已变','version':2})
        stale=client.get(path+'/research').json()[0]
        assert stale['result']=='' and stale['source_text']=='' and '不可用' in stale['error']
    finally:
        main.app.dependency_overrides.pop(main.db_session,None)
        engine.dispose()

def test_prd_independent_creation_and_word_export():
    from io import BytesIO
    from docx import Document
    from fastapi.testclient import TestClient
    from app.database import Base as CurrentBase
    engine=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool)
    CurrentBase.metadata.create_all(engine)
    def test_db():
        with Session(engine,expire_on_commit=False) as db:yield db
    main.app.dependency_overrides[main.db_session]=test_db
    try:
        client=TestClient(main.app)
        row=client.post('/api/v1/product-work',json={'title':'虚构PRD','kind':'prd'}).json()
        assert row['kind']=='prd' and row['project_id'] is None
        content='# 产品目标\n\n帮助学生查看到期提醒。\n\n| 角色 | 功能 |\n| --- | --- |\n| 学生 | 查看提醒 |'
        result=client.put('/api/v1/product-work/'+row['id'],json={'title':'虚构PRD','content':content,'version':1})
        assert result.status_code==200
        assert client.get('/api/v1/product-work/'+row['id']+'/export-word?version=1').status_code==409
        exported=client.get('/api/v1/product-work/'+row['id']+'/export-word?version=2')
        assert exported.status_code==200
        doc=Document(BytesIO(exported.content))
        assert any('产品目标' in p.text for p in doc.paragraphs)
        assert any('帮助学生查看到期提醒' in p.text for p in doc.paragraphs)
        assert doc.tables[0].cell(1,1).text=='查看提醒'
    finally:
        main.app.dependency_overrides.pop(main.db_session,None)
        engine.dispose()

def test_prd_section_edit_preserves_other_sections_and_rejects_stale(monkeypatch):
    from fastapi.testclient import TestClient
    from app.database import Base as CurrentBase
    import app.product_research as research
    engine=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool)
    CurrentBase.metadata.create_all(engine)
    def session():return Session(engine,expire_on_commit=False)
    def test_db():
        with session() as db:yield db
    monkeypatch.setattr(main,'model_key',lambda:'fake')
    monkeypatch.setattr(main,'model_settings',lambda db:(['deepseek-test'],1))
    monkeypatch.setattr(research,'model_key',lambda:'fake')
    monkeypatch.setattr(research,'model_settings',lambda db:(['deepseek-test'],1))
    monkeypatch.setattr(research,'SessionLocal',session)
    class Reply:
        def raise_for_status(self):pass
        def json(self):return {'choices':[{'finish_reason':'stop','message':{'content':'{"report":"修改后的背景"}'}}]}
    monkeypatch.setattr(research.httpx,'post',lambda *args,**kwargs:Reply())
    main.app.dependency_overrides[main.db_session]=test_db
    try:
        client=TestClient(main.app)
        row=client.post('/api/v1/product-work',json={'title':'虚构 PRD','kind':'prd'}).json()
        path='/api/v1/product-work/'+row['id']
        original='# 背景\n\n原背景\n\n# 目标\n\n目标保持不变\n'
        client.put(path,json={'title':'虚构 PRD','content':original,'version':1})
        assert [x['title'] for x in client.get(path+'/sections').json()]==['背景','目标']
        payload={'index':0,'instruction':'调整背景','version':2,'model':'deepseek-test','confirm_external_use':True}
        response=client.post(path+'/section-edits',json=payload);assert response.status_code==202
        task=client.get(path+'/section-edits').json()[0];assert task['status']=='completed'
        assert client.get(path).json()['content']==original
        accepted=client.post(path+'/section-edits/'+task['id']+'/decision',json={'decision':'accept'})
        assert accepted.status_code==200 and accepted.json()['version']==3
        assert accepted.json()['content']=='# 背景\n修改后的背景\n\n# 目标\n\n目标保持不变\n'
        assert client.get(path+'/revisions').json()[0]['content']==original
        payload['version']=3
        next_task=client.post(path+'/section-edits',json=payload).json()
        client.put(path,json={'title':'虚构 PRD','content':'# 背景\n人工修改\n\n# 目标\n原目标','version':3})
        stale=client.post(path+'/section-edits/'+next_task['id']+'/decision',json={'decision':'accept'})
        assert stale.status_code==409 and '人工修改' in client.get(path).json()['content']
    finally:
        main.app.dependency_overrides.pop(main.db_session,None)
        engine.dispose()

def test_prd_section_locator_ignores_code_and_rejects_unsafe_headings():
    from app.product_sections import headings, validate_replacement
    document='# 背景\n```markdown\n# 不是真正的章节\n```\n\n# 目标\n目标正文'
    assert [item['title'] for item in headings(document)]==['背景','目标']
    assert validate_replacement('新背景\n\n## 子标题\n解释',1)
    assert not validate_replacement('新背景\n\n# 越界章节',1)
    assert not validate_replacement('```\n未关闭',1)

def test_competitor_structured_comparison_history_and_status():
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from sqlalchemy.pool import StaticPool
    from app.database import Base as CurrentBase
    from app.security_context import actor_id, multiuser
    import app.main as main
    import pytest
    engine=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool);CurrentBase.metadata.create_all(engine)
    with Session(engine,expire_on_commit=False) as db:
        a=actor_id.set('local-user');m=multiuser.set(False)
        try:
            work=main.create_product_work(main.ProductWorkCreate(kind='competitor',title='对比'),db)
            body=main.CompetitorComparisonInput(version=1,competitors=[{'id':'one','name':'竞品一'},{'id':'two','name':'竞品二'}],observations=[{'id':'a','competitor_id':'one','dimension':'提醒','source_type':'inference','observation':'待核对'}])
            result=main.save_competitor_comparison(work.id,body,db)
            assert result.version==2 and len(main.get_competitor_comparison(work.id,db)['competitors'])==2
            with pytest.raises(main.ApiError) as stale:main.save_competitor_comparison(work.id,body,db)
            assert stale.value.status==409
            result=main.set_competitor_status(work.id,main.CompetitorStatusInput(version=2,status='completed'),db)
            assert result.status=='completed' and main.competitor_status_history(work.id,db)[0].new_status=='completed'
            assert main.get_competitor_comparison_history(work.id,db)[0]['version']==1
            actor_id.set('another')
            with pytest.raises(main.ApiError) as hidden:main.get_competitor_comparison(work.id,db)
            assert hidden.value.status==404
        finally:multiuser.reset(m);actor_id.reset(a)
    engine.dispose()

def test_competitor_selected_materials_only_and_export(monkeypatch):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from sqlalchemy.pool import StaticPool
    from fastapi import BackgroundTasks
    from app.database import Base as CurrentBase
    from app.security_context import actor_id, multiuser
    import app.main as main
    engine=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool);CurrentBase.metadata.create_all(engine)
    monkeypatch.setattr(main,'model_key',lambda:'fake');monkeypatch.setattr(main,'model_settings',lambda db:(['test-model'],1))
    with Session(engine,expire_on_commit=False) as db:
        a=actor_id.set('local-user');m=multiuser.set(False)
        try:
            work=main.create_product_work(main.ProductWorkCreate(kind='competitor',title='比较'),db)
            body=main.CompetitorComparisonInput(version=1,competitors=[{'id':'one','name':'选中产品'},{'id':'two','name':'未选产品'}],observations=[{'id':'a','competitor_id':'one','dimension':'流程','source_type':'inference','observation':'选中观察'},{'id':'b','competitor_id':'two','dimension':'流程','observation':'不应发送'}])
            main.save_competitor_comparison(work.id,body,db)
            task=main.start_product_research(work.id,main.ResearchInput(version=2,goal='比较',model='test-model',confirm_external_use=True,competitor_ids=['one']),BackgroundTasks(),db)
            assert '选中产品' in task.source_text and '选中观察' in task.source_text and '未选产品' not in task.source_text and '不应发送' not in task.source_text
            exported=main.export_product_work(work.id,db).body.decode()
            assert '选中产品' in exported and '未选产品' in exported and '待核查' in exported
        finally:multiuser.reset(m);actor_id.reset(a)
    engine.dispose()

def test_competitor_images_validate_content_and_preserve_versions():
    from io import BytesIO
    from PIL import Image
    from fastapi.testclient import TestClient
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from sqlalchemy.pool import StaticPool
    from app.database import Base as CurrentBase
    import app.main as main
    engine=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool);CurrentBase.metadata.create_all(engine)
    def dep():
        with Session(engine,expire_on_commit=False) as db:yield db
    main.app.dependency_overrides[main.db_session]=dep
    try:
        c=TestClient(main.app);work=c.post('/api/v1/product-work',json={'kind':'competitor','title':'截图测试'}).json();path='/api/v1/product-work/'+work['id']
        bad=c.post(path+'/images',data={'version':'1'},files={'file':('evil.png',b'<svg>invalid</svg>','image/png')})
        assert bad.status_code==422 and c.get(path+'/images').json()==[]
        image=BytesIO();Image.new('RGB',(8,8),'blue').save(image,'PNG')
        uploaded=c.post(path+'/images',data={'version':'1'},files={'file':('../../截图.png',image.getvalue(),'image/png')})
        assert uploaded.status_code==201 and uploaded.json()['version']==2
        row=c.get(path+'/images').json()[0];assert '/' not in row['filename'] and row['media_type']=='image/png'
        result=c.get(path+'/images/'+row['id']);assert result.status_code==200 and result.content==image.getvalue() and result.headers['x-content-type-options']=='nosniff'
        stale=c.post(path+'/images',data={'version':'1'},files={'file':('x.png',image.getvalue(),'image/png')})
        assert stale.status_code==409 and len(c.get(path+'/images').json())==1
    finally:main.app.dependency_overrides.pop(main.db_session,None);engine.dispose()

def test_competitor_image_permission_follows_project_membership():
    import asyncio
    from io import BytesIO
    from PIL import Image
    from fastapi import UploadFile
    engine=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool);Base.metadata.create_all(engine)
    with Session(engine,expire_on_commit=False) as db:
        owner=UserAccount(username='o',password_hash='x',active=True);member=UserAccount(username='m',password_hash='x',active=True)
        db.add_all([owner,member]);db.flush();project=Project(name='项目',background='');initialize_project_owner(db,project,owner.id);db.commit();add_member(db,project.id,owner.id,member.id)
        image=BytesIO();Image.new('RGB',(4,4),'blue').save(image,format='PNG');data=image.getvalue()
        with identity(owner.id):
            work=main.create_product_work(main.ProductWorkCreate(kind='competitor',title='对比',project_id=project.id),db)
            asyncio.run(main.upload_competitor_image(work.id,version=1,file=UploadFile(filename='a.png',file=BytesIO(data)),db=db))
            image_id=main.list_competitor_images(work.id,db)[0]['id']
        with identity(member.id):
            assert main.get_competitor_image(work.id,image_id,db).body==data
            with pytest.raises(main.ApiError) as forbidden:asyncio.run(main.upload_competitor_image(work.id,version=2,file=UploadFile(filename='a.png',file=BytesIO(data)),db=db))
            assert forbidden.value.status==403
        remove_member(db,project.id,owner.id,member.id)
        with identity(member.id):
            with pytest.raises(main.ApiError) as removed:main.get_competitor_image(work.id,image_id,db)
            assert removed.value.status==404
    engine.dispose()
