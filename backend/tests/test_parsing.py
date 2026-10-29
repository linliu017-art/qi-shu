import io
from docx import Document
import pytest
from app.parsing import ParseError, parse_content

def test_docx_reads_paragraphs_and_table():
    doc = Document()
    doc.add_paragraph('虚构客户需求')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '项目'
    table.cell(0, 1).text = '内容'
    table.cell(1, 0).text = '管理'
    table.cell(1, 1).text = '台账'
    f = io.BytesIO()
    doc.save(f)
    text, warning = parse_content('资料.docx', f.getvalue())
    assert '虚构客户需求' in text and '管理 | 台账' in text
    assert '复杂' in warning

def test_reject_non_docx_and_scanned_pdf():
    with pytest.raises(ParseError):
        parse_content('假文件.docx', b'not a zip')
    with pytest.raises(ParseError):
        parse_content('假文件.pdf', b'not pdf')

def test_pdf_reads_text_and_simple_table():
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    pdf = canvas.Canvas(buf, pagesize=(300, 300))
    pdf.drawString(40, 270, 'Interview notes')
    for x in (40, 140, 260): pdf.line(x, 120, x, 220)
    for y in (120, 170, 220): pdf.line(40, y, 260, y)
    pdf.drawString(45, 190, 'Topic')
    pdf.drawString(145, 190, 'Requirement')
    pdf.drawString(45, 140, 'Task')
    pdf.drawString(145, 140, 'Tracking')
    pdf.save()
    text, warning = parse_content('notes.pdf', buf.getvalue())
    assert '[第1页]' in text
    assert 'Topic | Requirement' in text and 'Task | Tracking' in text
    assert '复杂表格' in warning

def test_quote_alignment_ignores_only_whitespace():
    from app.analysis import align_quote
    assert align_quote('负责人 确认完成', '由负责人\n确认完成。') == '负责人\n确认完成'
    assert align_quote('负责人批准完成', '由负责人\n确认完成。') is None
    assert align_quote('  \n ', '任意内容') is None

def test_docx_reports_location_of_merged_table():
    doc = Document()
    doc.add_paragraph('项目背景')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '字段'
    table.cell(0, 1).text = '说明'
    table.cell(1, 0).merge(table.cell(1, 1)).text = '待确认'
    content = io.BytesIO()
    doc.save(content)
    text, warning = parse_content('虚构.docx', content.getvalue())
    assert '[第1段]' in text and '[第1表]' in text
    assert '第1表含合并或嵌套单元格' in warning


def test_pdf_reports_image_page_with_readable_text():
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader
    from PIL import Image
    image_data = io.BytesIO()
    Image.new('RGB', (4, 4), color='white').save(image_data, format='PNG')
    image_data.seek(0)
    content = io.BytesIO()
    pdf = canvas.Canvas(content)
    pdf.drawString(30, 700, 'Text present on image page')
    pdf.drawImage(ImageReader(image_data), 30, 650, width=10, height=10)
    pdf.save()
    text, warning = parse_content('虚构.pdf', content.getvalue())
    assert 'Text present on image page' in text
    assert '检测到图片的页面（图片内容未识别）：1' in warning
