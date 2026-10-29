from contextlib import contextmanager
from io import BytesIO
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
import app.main as main
from app.database import Base
from app.membership import initialize_project_owner, add_member, remove_member
from app.models import Project, UserAccount, ManualRequirement
from app.security_context import actor_id, multiuser

@contextmanager
def identity(user):
    a=actor_id.set(user);m=multiuser.set(True)
    try:yield
    finally:multiuser.reset(m);actor_id.reset(a)

def test_project_members_can_edit_table_but_outsiders_cannot():
    engine=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine,expire_on_commit=False) as db:
        owner=UserAccount(username='owner',password_hash='test',active=True)
        member=UserAccount(username='member',password_hash='test',active=True)
        db.add_all([owner,member]);db.flush()
        project=Project(name='虚构项目',background='');initialize_project_owner(db,project,owner.id);db.commit()
        add_member(db,project.id,owner.id,member.id)
        with identity(owner.id):
            row=main.create_product_requirement(main.ProductRequirementCreate(title='需求A',project_id=project.id),db)
            private=main.create_product_requirement(main.ProductRequirementCreate(title='个人需求'),db)
        with identity(member.id):
            changed=main.edit_product_requirement(row.id,main.ProductRequirementEdit(title='需求A',status='doing',version=1),db)
            assert changed.status=='doing' and changed.version==2
            with pytest.raises(main.ApiError) as hidden:main.get_product_requirement(private.id,db)
            assert hidden.value.status==404
            with pytest.raises(main.ApiError) as stale:main.edit_product_requirement(row.id,main.ProductRequirementEdit(title='冲突',version=1),db)
            assert stale.value.status==409
        remove_member(db,project.id,owner.id,member.id)
        with identity(member.id):
            with pytest.raises(main.ApiError) as removed:main.get_product_requirement(row.id,db)
            assert removed.value.status==404
    engine.dispose()

def test_import_snapshot_and_excel_is_not_formula():
    from app.database import Base as CurrentBase
    engine=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool)
    CurrentBase.metadata.create_all(engine)
    def db_dep():
        with Session(engine,expire_on_commit=False) as db:yield db
    main.app.dependency_overrides[main.db_session]=db_dep
    try:
        c=TestClient(main.app)
        pid=c.post('/api/v1/projects',json={'name':'虚构项目'}).json()['id']
        with Session(engine,expire_on_commit=False) as db:
            source=main.ManualRequirement(project_id=pid,text='原始需求',original_text='原始需求',status='confirmed',version=1,created_by='local-user')
            db.add(source);db.commit();sid=source.id
        imported=c.post('/api/v1/product-requirements/import',json={'project_id':pid,'source_type':'manual_requirement','source_id':sid})
        assert imported.status_code==201 and imported.json()['source_snapshot']=='原始需求'
        with Session(engine) as db:
            source=db.get(main.ManualRequirement,sid);source.text='已修改';source.version=2;db.commit()
        result=c.get('/api/v1/product-requirements/'+imported.json()['id']).json()
        assert result['description']=='原始需求' and result['source_state']=='changed'
        c.post('/api/v1/product-requirements',json={'project_id':pid,'title':'=1+1','description':'虚构描述'})
        exported=c.get('/api/v1/product-requirements-export',params={'project_id':pid})
        assert exported.status_code==200
        sheet=load_workbook(BytesIO(exported.content),data_only=False).active
        cell=next(cell for row in sheet.iter_rows() for cell in row if cell.value=='=1+1')
        assert cell.data_type=='s'
    finally:
        main.app.dependency_overrides.pop(main.db_session,None);engine.dispose()

def test_prd_candidates_require_consent_accept_once_and_preserve_progress(monkeypatch):
    from app.database import Base as CurrentBase
    import app.product_research as research
    engine=create_engine('sqlite://',connect_args={'check_same_thread':False},poolclass=StaticPool)
    CurrentBase.metadata.create_all(engine)
    def session():return Session(engine,expire_on_commit=False)
    def db_dep():
        with session() as db:yield db
    monkeypatch.setattr(main,'model_key',lambda:'fake')
    monkeypatch.setattr(main,'model_settings',lambda db:(['deepseek-test'],1))
    monkeypatch.setattr(research,'model_key',lambda:'fake')
    monkeypatch.setattr(research,'model_settings',lambda db:(['deepseek-test'],1))
    monkeypatch.setattr(research,'SessionLocal',session)
    class Reply:
        def raise_for_status(self):pass
        def json(self):return {'choices':[{'finish_reason':'stop','message':{'content':'{"requirements":[{"title":"归还提醒","description":"提醒学生归还图书","quote":"提醒学生归还图书"}]}'}}]}
    monkeypatch.setattr(research.httpx,'post',lambda *args,**kwargs:Reply())
    main.app.dependency_overrides[main.db_session]=db_dep
    try:
        c=TestClient(main.app)
        prd=c.post('/api/v1/product-work',json={'title':'虚构PRD','kind':'prd'}).json();path='/api/v1/product-work/'+prd['id']
        c.put(path,json={'title':'虚构PRD','content':'# 需求\n提醒学生归还图书','version':1})
        payload={'version':2,'model':'deepseek-test'}
        assert c.post(path+'/requirement-candidates',json=payload).status_code==400
        payload['confirm_external_use']=True
        started=c.post(path+'/requirement-candidates',json=payload);assert started.status_code==202
        task=c.get(path+'/requirement-candidates').json()[0]
        assert task['status']=='completed'
        assert c.get('/api/v1/product-requirements').json()==[]
        body={'candidates':[{'index':0,'title':'归还提醒','description':'提醒学生归还图书'}]}
        accepted=c.post(path+'/requirement-candidates/'+task['id']+'/accept',json=body)
        assert accepted.status_code==200 and len(accepted.json())==1
        req=accepted.json()[0]
        assert req['source_type']=='prd' and req['source_snapshot']=='提醒学生归还图书'
        assert c.post(path+'/requirement-candidates/'+task['id']+'/accept',json=body).status_code==409
        updated=c.put('/api/v1/product-requirements/'+req['id'],json={'title':req['title'],'description':req['description'],'status':'doing','version':1})
        assert updated.status_code==200
        c.put(path,json={'title':'虚构PRD','content':'# 新需求\n其他内容','version':2})
        saved=c.get('/api/v1/product-requirements/'+req['id']).json()
        assert saved['source_state']=='changed' and saved['status']=='doing'
    finally:
        main.app.dependency_overrides.pop(main.db_session,None);engine.dispose()
